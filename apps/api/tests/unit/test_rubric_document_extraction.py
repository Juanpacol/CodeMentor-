"""Tests unitarios de `modules/rubrics/document_extraction.py` — solo la
extracción de texto crudo, sin tocar la DB ni el harness de IA."""

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from logica.core.errors import ValidationDomainError
from logica.modules.rubrics import document_extraction


def _build_pdf_bytes(text: str) -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _build_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row_values in table_rows:
            row = table.add_row()
            for cell, value in zip(row.cells, row_values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_text_from_docx_incluye_parrafos_y_tablas() -> None:
    raw = _build_docx_bytes(
        ["Temario del curso"],
        table_rows=[["Unidad", "Tema"], ["1", "Estructuras condicionales"]],
    )
    text = document_extraction.extract_text_from_docx(raw)
    assert "Temario del curso" in text
    assert "Estructuras condicionales" in text


def test_extract_text_from_docx_vacio_lanza_error() -> None:
    raw = _build_docx_bytes([])
    with pytest.raises(ValidationDomainError):
        document_extraction.extract_text_from_docx(raw)


def test_extract_text_from_docx_corrupto_lanza_error() -> None:
    with pytest.raises(ValidationDomainError):
        document_extraction.extract_text_from_docx(b"esto no es un docx valido")


def test_extract_text_from_pdf_sin_texto_lanza_error() -> None:
    raw = _build_pdf_bytes("")
    with pytest.raises(ValidationDomainError):
        document_extraction.extract_text_from_pdf(raw)


def test_extract_text_from_pdf_corrupto_lanza_error() -> None:
    with pytest.raises(ValidationDomainError):
        document_extraction.extract_text_from_pdf(b"esto no es un pdf valido")


def test_truncate_recorta_texto_largo() -> None:
    long_paragraph = "Tema repetido. " * 3000
    raw = _build_docx_bytes([long_paragraph])
    text = document_extraction.extract_text_from_docx(raw)
    assert len(text) <= document_extraction._MAX_EXTRACTED_CHARS
