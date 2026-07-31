"""Extracción de texto crudo de una rúbrica institucional en PDF/DOCX (§ubir
rúbrica). Solo extrae texto plano — la interpretación semántica (¿cuáles son
los temas? ¿qué nivel tienen?) la hace el LLM en `rubric_topic_extractor.py`,
porque el formato institucional es demasiado heterogéneo (tablas, unidades,
niveles) para un parser de tablas hecho a mano."""

import io

import structlog
from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from logica.core.errors import ValidationDomainError

logger = structlog.get_logger()

# 5 MB: generoso para un temario institucional (típicamente unas pocas
# páginas) pero acotado para no cargar un archivo gigante a memoria.
MAX_FILE_BYTES = 5_000_000
ALLOWED_EXTENSIONS = (".pdf", ".docx")
ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

# Un temario de 80 páginas no aporta más que uno de 10 para este prompt y
# infla el costo/latencia de la llamada al modelo; se corta antes de mandarlo.
_MAX_EXTRACTED_CHARS = 20_000


def _truncate(text: str) -> str:
    if len(text) <= _MAX_EXTRACTED_CHARS:
        return text
    logger.warning("rubric_document_text_truncated", original_chars=len(text))
    return text[:_MAX_EXTRACTED_CHARS]


def extract_text_from_pdf(raw: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, ValueError) as exc:
        raise ValidationDomainError(
            "No se pudo leer el documento",
            hint="Verifica que el PDF no esté dañado o protegido con contraseña.",
        ) from exc
    return _require_readable_text(text)


def extract_text_from_docx(raw: bytes) -> str:
    try:
        document = Document(io.BytesIO(raw))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = " | ".join(cell.text.strip() for cell in row.cells)
                if cells.strip(" |"):
                    parts.append(cells)
        text = "\n".join(parts)
    except Exception as exc:  # pypdf/python-docx no exponen una jerarquía común
        raise ValidationDomainError(
            "No se pudo leer el documento",
            hint="Verifica que el archivo sea un .docx válido y no esté dañado.",
        ) from exc
    return _require_readable_text(text)


def _require_readable_text(text: str) -> str:
    if len(text.strip()) < 20:
        raise ValidationDomainError(
            "El documento no tiene texto legible",
            hint=(
                "Si es un PDF escaneado (una imagen del papel), esta función no puede "
                "leerlo — todavía no soportamos reconocimiento óptico de caracteres."
            ),
        )
    return _truncate(text)
